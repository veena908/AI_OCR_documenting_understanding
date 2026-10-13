import os
import uuid

from docx import Document
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

OUTPUT_FOLDER = "outputs"


def _output_path(ext: str) -> str:
    os.makedirs(OUTPUT_FOLDER, exist_ok=True)
    return os.path.join(OUTPUT_FOLDER, f"export_{uuid.uuid4().hex}.{ext}")


def _structured_lines(structured_info: dict) -> list[str]:
    lines = []

    for key, values in (structured_info or {}).items():
        if values:
            lines.append(f"{key.capitalize()}: {', '.join(values)}")

    return lines


def export_pdf(text: str, structured_info: dict, document_type: str) -> str:
    path = _output_path("pdf")
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(path, pagesize=A4)

    story = [
        Paragraph("AI OCR Document Export", styles["Title"]),
        Spacer(1, 0.2 * inch),
        Paragraph(f"Document Type: {document_type}", styles["Heading2"]),
        Spacer(1, 0.2 * inch),
    ]

    for line in _structured_lines(structured_info):
        story.append(Paragraph(line, styles["Normal"]))

    story.append(Spacer(1, 0.2 * inch))
    story.append(Paragraph("Extracted Text", styles["Heading2"]))

    for paragraph in text.split("\n"):
        story.append(Paragraph(paragraph if paragraph.strip() else "&nbsp;", styles["Normal"]))

    doc.build(story)

    return path


def export_docx(text: str, structured_info: dict, document_type: str) -> str:
    path = _output_path("docx")
    doc = Document()

    doc.add_heading("AI OCR Document Export", level=0)
    doc.add_heading(f"Document Type: {document_type}", level=2)

    for line in _structured_lines(structured_info):
        doc.add_paragraph(line)

    doc.add_heading("Extracted Text", level=2)

    for paragraph in text.split("\n"):
        doc.add_paragraph(paragraph)

    doc.save(path)

    return path


def export_xlsx(text: str, structured_info: dict, document_type: str) -> str:
    path = _output_path("xlsx")
    wb = Workbook()

    info_sheet = wb.active
    info_sheet.title = "Document Info"
    info_sheet.append(["Document Type", document_type])
    info_sheet.append([])
    info_sheet.append(["Field", "Value"])

    for key, values in (structured_info or {}).items():
        info_sheet.append([key.capitalize(), ", ".join(values) if values else ""])

    text_sheet = wb.create_sheet("Extracted Text")

    for line_number, line in enumerate(text.split("\n"), start=1):
        text_sheet.cell(row=line_number, column=1, value=line)

    wb.save(path)

    return path
